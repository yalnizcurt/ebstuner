"""Bounded source-spec extraction and declarative, byte-preserving DAT validation."""
from __future__ import annotations

import copy
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from engine import DATE_FORMATS, check_encoding, parse_lines
from compiler import parse_condition
from spreadsheets import MAX_CELLS, MAX_ROWS, MAX_UNCOMPRESSED, cell_text

DISPLAY_COLUMNS = [
    {'key': 'field_name', 'label': 'Field Name'},
    {'key': 'category_functional_group', 'label': 'Category / Functional Group'},
    {'key': 'mapping_type', 'label': 'Mapping Type'},
    {'key': 'start_pos', 'label': 'Start Pos'},
    {'key': 'end_pos', 'label': 'End Pos'},
    {'key': 'length', 'label': 'Length'},
    {'key': 'datatype_conversion', 'label': 'Data Type / Oracle Conversion Function'},
    {'key': 'padding_alignment', 'label': 'Padding / Alignment Rule'},
    {'key': 'raw_sample', 'label': 'Raw Sample Value in DAT File (Quoted)'},
    {'key': 'cleansed_output', 'label': 'Cleansed Output in Oracle Staging Table'},
]
MASTER_COLUMNS = [{'key': 'id', 'label': 'Field ID'}, *DISPLAY_COLUMNS, *[
    {'key': 'sequence', 'label': 'Sequence'}, {'key': 'record_type', 'label': 'Record Type'},
    {'key': 'mandatory', 'label': 'Mandatory'}, {'key': 'alignment', 'label': 'Alignment'},
    {'key': 'padding_byte', 'label': 'Padding Byte'}, {'key': 'padding_direction', 'label': 'Padding Direction'},
    {'key': 'format', 'label': 'Format'}, {'key': 'allowed_values', 'label': 'Allowed Values'},
    {'key': 'conditional_rules', 'label': 'Conditional Rules'}, {'key': 'description', 'label': 'Description'},
    {'key': 'remarks', 'label': 'Remarks'}, {'key': 'source_reference', 'label': 'Source Reference'},
]]
COLUMN_KEYS = [c['key'] for c in MASTER_COLUMNS]
MAX_BYTES, MAX_FIELDS, MAX_RECORD_WIDTH = 12_000_000, 4000, 1_000_000


class SpecError(ValueError):
    def __init__(self, message, issues=None):
        super().__init__(message)
        self.issues = issues or [{'code': 'INVALID_SPEC', 'message': message, 'severity': 'error'}]


def issue(code, message, field_id=None, record=None, severity='error'):
    result = {'code': code, 'message': message, 'severity': severity}
    if field_id is not None: result['field_id'] = field_id
    if record is not None: result['record'] = record
    return result


def _header(value):
    return re.sub(r'[^a-z0-9]+', '', str(value).casefold())


HEADER_KEYS = {_header(c['label']): c['key'] for c in MASTER_COLUMNS}
HEADER_KEYS.update({_header(k): k for k in COLUMN_KEYS})
HEADER_KEYS.update({'category': 'category_functional_group', 'functionalgroup': 'category_functional_group',
                    'start': 'start_pos', 'startposition': 'start_pos', 'end': 'end_pos', 'endposition': 'end_pos',
                    'datatype': 'datatype_conversion', 'oracleconversionfunction': 'datatype_conversion',
                    'padding': 'padding_alignment',
                    'rawsamplevalue': 'raw_sample', 'cleansedoutput': 'cleansed_output',
                    'required': 'mandatory', 'mandatoryfield': 'mandatory', 'fieldidentifier': 'id'})


def _md_cells(line):
    # Escaped pipes belong to the cell; pipes in code spans are not table separators.
    result, current, escaped, code = [], [], False, False
    for char in line.strip():
        if escaped:
            current.append(char if char == '|' else '\\' + char); escaped = False
        elif char == '\\': escaped = True
        elif char == '`': code = not code; current.append(char)
        elif char == '|' and not code: result.append(''.join(current).strip()); current = []
        else: current.append(char)
    if escaped: current.append('\\')
    result.append(''.join(current).strip())
    if line.strip().startswith('|'): result = result[1:]
    if line.strip().endswith('|'): result = result[:-1]
    return result


def _source_tables(name, raw):
    extension = Path(name).suffix.lower()
    warnings = []
    if extension == '.csv':
        text = raw.decode('utf-8-sig', 'strict')
        try: dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=''), dialect)
        rows = []
        for row in reader:
            rows.append({'row': reader.line_num, 'values': row,
                         'references': [f'CSV!{reader.line_num}:{i + 1}' for i in range(len(row))]})
            if len(rows) > MAX_ROWS: raise SpecError('Specification exceeds 20,000 rows')
        return [('CSV', rows)], warnings, text
    if extension in ('.md', '.markdown'):
        text = raw.decode('utf-8-sig', 'strict')
        tables, rows, section, fenced = [], [], 'Markdown', False
        for number, line in enumerate(text.splitlines(), 1):
            if line.lstrip().startswith(('```', '~~~')): fenced = not fenced; continue
            if fenced: continue
            if line.lstrip().startswith('#'):
                if rows: tables.append((section, rows)); rows = []
                section = line.strip().lstrip('#').strip() or f'Section {number}'
            elif '|' in line:
                values = _md_cells(line)
                if all(re.fullmatch(r':?-{3,}:?', v.replace(' ', '')) for v in values): continue
                rows.append({'row': number, 'values': values,
                             'references': [f'Markdown!line{number}:column{i + 1}' for i in range(len(values))]})
            elif line.strip():
                rows.append({'row': number, 'values': [line], 'references': [f'Markdown!line{number}'], 'metadata': True})
            if sum(len(r) for _, r in tables) + len(rows) > MAX_ROWS: raise SpecError('Specification exceeds 20,000 rows')
        if rows: tables.append((section, rows))
        return tables, warnings, text
    if extension != '.xlsx': raise SpecError('Specification must be CSV, XLSX or Markdown')
    from openpyxl import load_workbook
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > MAX_UNCOMPRESSED:
            raise SpecError('XLSX exceeds the bounded decompression limit')
    workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
    tables, cells = [], 0
    try:
        if len(workbook.worksheets) > 30: raise SpecError('At most 30 worksheets are supported')
        for sheet in workbook.worksheets:
            rows = []
            for number, row in enumerate(sheet.iter_rows(), 1):
                if number > MAX_ROWS or len(row) > 200: raise SpecError('Worksheet exceeds row/column limits')
                cells += len(row)
                if cells > MAX_CELLS: raise SpecError('Specification exceeds 400,000 cells')
                values, refs = [], []
                for index, cell in enumerate(row, 1):
                    reference = f'{sheet.title}!{getattr(cell, "coordinate", str(number) + ":" + str(index))}'
                    value = cell_text(cell.value, cell.number_format)
                    if cell.data_type == 'f':
                        warnings.append(issue('FORMULA', f'{reference}: formula retained as source text and never executed'))
                    if isinstance(cell.value, float):
                        warnings.append(issue('EXCEL_PRECISION', f'{reference}: stored Excel numeric precision applies', severity='warning'))
                    values.append(value); refs.append(reference)
                rows.append({'row': number, 'values': values, 'references': refs})
            tables.append((sheet.title, rows))
    finally: workbook.close()
    return tables, warnings, None


def extract_spec(name, raw, column_mapping=None):
    """Preserve every source column and its location; no model or formula execution."""
    if not isinstance(raw, bytes) or not raw or len(raw) > MAX_BYTES:
        raise SpecError('Specification must contain 1 byte to 12 MB')
    try: tables, warnings, source_text = _source_tables(name, raw)
    except (UnicodeError, csv.Error, zipfile.BadZipFile) as exc: raise SpecError('Unreadable specification document') from exc
    document = {'filename': Path(name).name, 'sha256': hashlib.sha256(raw).hexdigest(),
                'format': Path(name).suffix.lower().lstrip('.'), 'columns': copy.deepcopy(DISPLAY_COLUMNS),
                'sections': [], 'fields': [], 'issues': warnings, 'options': {}}
    if source_text is not None: document['source_text'] = source_text
    cells = 0
    for table_number, (name, rows) in enumerate(tables, 1):
        section = {'id': f'section-{table_number}', 'name': name, 'headers': [], 'raw_rows': rows, 'metadata': []}
        document['sections'].append(section)
        mapping = None
        for row in rows:
            values = row['values']; cells += len(values)
            if cells > MAX_CELLS or len(values) > 200: raise SpecError('Specification exceeds cell/column limits')
            if row.get('metadata'):
                section['metadata'].append(row); continue
            guessed = {i: HEADER_KEYS[_header(v)] for i, v in enumerate(values) if _header(v) in HEADER_KEYS}
            if column_mapping:
                guessed = {i: column_mapping.get(v, column_mapping.get(str(i), guessed.get(i))) for i, v in enumerate(values)}
                guessed = {i: key for i, key in guessed.items() if key in COLUMN_KEYS}
            if 'field_name' in guessed.values() and len(guessed) >= 2:
                if len(set(guessed.values())) != len(guessed):
                    warnings.append(issue('DUPLICATE_COLUMN', f'{name} row {row["row"]}: duplicate specification column'))
                mapping = guessed; section['headers'] = values; section['header_row'] = row['row']
                missing = [c['label'] for c in DISPLAY_COLUMNS if c['key'] not in guessed.values()]
                if missing: warnings.append(issue('MISSING_COLUMNS', f'{name}: missing source columns: ' + ', '.join(missing), severity='warning'))
                continue
            if mapping is None:
                if any(str(v).strip() for v in values): section['metadata'].append(row)
                continue
            if not any(str(v).strip() for v in values): continue
            field = {key: '' for key in COLUMN_KEYS}
            for index, key in mapping.items(): field[key] = values[index] if index < len(values) else ''
            if not field['field_name'].strip():
                warnings.append(issue('UNMAPPED_ROW', f'{name} row {row["row"]}: nonempty source row has no field name'))
                continue
            field['id'] = field['id'] or f'{section["id"]}-row-{row["row"]}'
            field['sequence'] = field['sequence'] or str(len(document['fields']) + 1)
            field['mandatory'] = field['mandatory'].upper() or 'UNKNOWN'
            field['source'] = {'section': section['id'], 'sheet': name, 'row': row['row'],
                               'cells': {key: {'value': values[i] if i < len(values) else '',
                                               'reference': row['references'][i] if i < len(row['references']) else f'{name}!{row["row"]}:{i + 1}'} for i, key in mapping.items()},
                               'extra_columns': {section['headers'][i] if i < len(section['headers']) else str(i + 1): value
                                                 for i, value in enumerate(values) if i not in mapping}}
            field['rules'] = infer_rules(field)
            document['fields'].append(field)
            if len(document['fields']) > MAX_FIELDS: raise SpecError('Specification exceeds 4,000 fields')
    if not document['fields']: warnings.append(issue('NO_FIELDS', 'No recognized field table; review the retained source headings and map its columns'))
    return document


def _integer(value):
    if isinstance(value, bool): return None
    if isinstance(value, int): return value
    value = str(value or '').strip()
    return int(value) if re.fullmatch(r'[0-9]{1,7}', value) else None


def _type_rule(text):
    value = text.strip(); upper = value.upper()
    char_trim = re.fullmatch(r'(?:CHAR|VARCHAR2?)(?:\s*\(\d+\))?\s*/\s*(TRIM|RTRIM|LTRIM)(?:\(\s*:[A-Z0-9_]+\s*\))?', upper)
    if char_trim: return {'datatype': 'char', 'conversion': char_trim.group(1).lower()}
    if re.fullmatch(r'(?:CHAR|VARCHAR2?)(?:\s*\(\d+\))?\s*/\s*TO_NUMBER(?:\(\s*:[A-Z0-9_]+\s*\))?', upper):
        return {'datatype': 'decimal', 'conversion': 'to_number'}
    if re.fullmatch(r'(?:CHAR|VARCHAR2?|TEXT|STRING)(?:\s*\(\d+\))?', upper): return {'datatype': 'char', 'conversion': 'identity'}
    if re.fullmatch(r'(?:INTEGER|INT)\s+EXTERNAL(?:\s*\(\d+\))?', upper): return {'datatype': 'integer', 'conversion': 'to_number'}
    number = re.fullmatch(r'(?:NUMBER|NUMERIC)(?:\s*\(\s*(\d+)\s*(?:,\s*([+-]?\d+)\s*)?\))?', upper)
    if number:
        rule = {'datatype': 'decimal', 'conversion': 'to_number'}
        if number.group(1): rule.update(precision=int(number.group(1)), scale=int(number.group(2) or 0))
        return rule
    if re.fullmatch(r'DECIMAL\s+EXTERNAL(?:\s*\(\d+\))?', upper): return {'datatype': 'decimal', 'conversion': 'to_number'}
    if re.fullmatch(r'FLOAT\s+EXTERNAL(?:\s*\(\d+\))?', upper): return {'datatype': 'float', 'conversion': 'to_number'}
    date = re.fullmatch(r'DATE(?:\s*\(\d+\))?\s+[\'"]([^\'"]+)[\'"]', value, re.I)
    if date: return {'datatype': 'date', 'conversion': 'to_date', 'date_format': date.group(1).upper()}
    trim = re.fullmatch(r'(TRIM|RTRIM|LTRIM)\(\s*:[A-Z0-9_]+\s*\)', upper)
    if trim: return {'datatype': 'char', 'conversion': trim.group(1).lower()}
    if re.fullmatch(r'TO_NUMBER\(\s*(?::[A-Z0-9_]+|TRIM\(\s*:[A-Z0-9_]+\s*\))\s*\)', upper): return {'datatype': 'decimal', 'conversion': 'to_number'}
    date = re.fullmatch(r'TO_DATE\(\s*:[A-Z0-9_]+\s*,\s*[\'"]([^\'"]+)[\'"]\s*\)', value, re.I)
    if date: return {'datatype': 'date', 'conversion': 'to_date', 'date_format': date.group(1).upper()}
    return {'datatype': None, 'conversion': None, 'unsupported_conversion': value or 'Missing datatype/conversion'}


def infer_rules(field):
    mapping = str(field.get('mapping_type', '')).strip().upper()
    kind = {'POSITION': 'physical', 'PHYSICAL': 'physical', 'FIELD': 'physical', 'DATA': 'physical',
            'FILLER': 'filler', 'RESERVED': 'filler', 'CONSTANT': 'constant',
            'SYSTEM': 'system', 'EXPRESSION': 'system', 'SPECIAL/EXPR': 'system', 'GENERATED': 'system'}.get(mapping)
    rules = {'mapping_type': kind, 'start': _integer(field.get('start_pos')), 'end': _integer(field.get('end_pos')),
             'length': _integer(field.get('length')), 'required': None, 'layout_id': 'primary'}
    if kind == 'system': return {**rules, 'datatype': 'system', 'conversion': 'database', 'required': False, 'padding': {'side': 'none', 'char': ' '}}
    if kind == 'constant':
        value = str(field.get('datatype_conversion', '')).strip()
        quoted = re.fullmatch(r"'((?:[^']|'')*)'", value)
        return {**rules, 'datatype': 'char', 'conversion': 'identity', 'required': False,
                'constant': quoted.group(1).replace("''", "'") if quoted else None, 'padding': {'side': 'none', 'char': ' '}}
    rules.update(_type_rule(str(field.get('datatype_conversion', ''))))
    text = str(field.get('padding_alignment', '')).strip().lower()
    char = '0' if 'zero' in text else ' '
    if 'trailing' in text or 'left align' in text: side = 'right'
    elif 'leading' in text or 'right align' in text: side = 'left'
    elif 'exact width' in text or 'fixed text' in text or 'fixed date' in text or text in ('none', 'no padding'): side = 'none'
    else: side = None
    rules['padding'] = {'side': side, 'char': char}
    if kind == 'filler': rules.update(required=False, filler_value='')
    return rules


def _routing_conditions(expressions, fields):
    when, by_name = [], {}
    for field in fields:
        rules = field.get('rules') or infer_rules(field)
        if rules.get('start') and rules.get('end'):
            by_name.setdefault(field['field_name'].strip().upper(), set()).add((rules['start'], rules['end']))

    def convert(node):
        if node[0] == 'and':
            for child in node[1]: convert(child)
            return
        if node[0] != 'compare' or node[2] != '=' or node[3][0] != 'literal':
            raise ValueError('Only explicit AND equality routing is supported')
        operand = node[1]
        if operand[0] == 'field':
            spans = by_name.get(operand[1], set())
            if len(spans) != 1: raise ValueError('Routing field is missing or has ambiguous positions: ' + operand[1])
            start, end = next(iter(spans)); trim = True
        elif operand[0] == 'range': start, end, trim = operand[1], operand[2], False
        else: raise ValueError('Routing must reference a known source field or an explicit byte range')
        condition = {'start': start, 'end': end, 'equals': node[3][1], 'trim': trim}
        if condition not in when: when.append(condition)

    for expression in expressions: convert(parse_condition(expression))
    return when


def infer_layouts(document, fields=None):
    """Translate only explicitly labelled source conditions; never guess a record identifier."""
    fields = fields if fields is not None else document.get('fields', [])
    layouts, findings = [], []
    sections = document.get('sections', [])
    if not sections: return {'layouts': [], 'issues': []}
    for section in sections:
        selected = [f for f in fields if f.get('source', {}).get('section') == section['id']]
        if not selected: continue
        expressions = []
        for row in section.get('metadata', []):
            text = ' '.join(str(value) for value in row.get('values', []) if str(value).strip()).strip()
            match = re.match(r'(?:CTL\s+)?(?:ROUTING\s+CONDITION|RECORD\s+CONDITION|WHEN)\s*:\s*(.+)$', text, re.I)
            if match: expressions.append((match.group(1), row.get('references', [])))
        try:
            when = _routing_conditions([expression for expression, _ in expressions], selected)
        except ValueError as exc:
            findings.append(issue('SOURCE_ROUTING', f'{section["name"]}: source routing needs explicit review: {exc}'))
            when = []
        layouts.append({'id': section['id'], 'name': section['name'], 'field_ids': [f['id'] for f in selected],
                        'record_length': max((f.get('rules', {}).get('end') or 0 for f in selected), default=0),
                        'when': when, 'source_conditions': [{'text': text, 'references': refs} for text, refs in expressions]})
    return {'layouts': layouts, 'issues': findings}


def import_document(document, options=None):
    """Keep source conditions and uncompiled source text visible at approval."""
    fields = copy.deepcopy(document.get('fields', []))
    inferred = infer_layouts(document, fields)
    explicit = any(layout.get('source_conditions') for layout in inferred['layouts'])
    layouts = inferred['layouts'] if explicit else []
    if explicit:
        for layout in layouts:
            identifiers = {field['record_type'].strip() for field in fields
                           if field['id'] in layout['field_ids'] and field.get('record_type', '').strip()}
            if len(identifiers) == 1: layout['id'] = next(iter(identifiers))
    requirements = []
    for section in document.get('sections', []):
        for row in section.get('metadata', []):
            text = ' '.join(str(value) for value in row.get('values', []) if str(value).strip()).strip()
            if not text or re.match(r'(?:CTL\s+)?(?:ROUTING\s+CONDITION|RECORD\s+CONDITION|WHEN)\s*:', text, re.I): continue
            reference = ', '.join(row.get('references', []))
            requirements.append({'id': f"metadata:{section['id']}:{row['row']}", 'reference': reference,
                                 'text': text, 'status': 'unresolved'})
    for field in fields:
        for column, value in field.get('source', {}).get('extra_columns', {}).items():
            if str(value).strip():
                requirements.append({'id': f"extra:{field['id']}:{column}",
                                     'reference': f"{field['source'].get('sheet', '')}!{field['source'].get('row', '')}:{column}",
                                     'text': f'{column}: {value}', 'status': 'unresolved'})
    return normalize_spec({'fields': fields, 'layouts': layouts, 'options': options or {},
                           'source_issues': [*document.get('issues', []), *inferred['issues']],
                           'source_requirements': requirements, 'file_rules': []})


def spec_fingerprint(spec):
    """Bind human decisions to executable rules, independent of UI review state."""
    normalized = normalize_spec(spec)
    options = {key: value for key, value in normalized['options'].items()
               if key not in ('confirmed_records', 'suspect_records', 'mapping_resolutions', 'detail_record', 'record_number')}
    contract = {'fields': [{'id': field['id'], 'rules': field['rules']} for field in normalized['fields']],
                'layouts': normalized['layouts'], 'options': options, 'file_rules': normalized.get('file_rules', [])}
    return hashlib.sha256(json.dumps(contract, sort_keys=True, separators=(',', ':'), ensure_ascii=False).encode()).hexdigest()


def _layouts(spec):
    if spec.get('layouts'): return spec['layouts']
    groups = {}
    for field in spec.get('fields', []):
        identifier = field.get('record_type', '').strip() or field.get('rules', {}).get('layout_id', 'primary')
        groups.setdefault(identifier, []).append(field)
    layouts = []
    for name, fields in groups.items():
        sources = {f.get('conditional_rules', '').strip() for f in fields if f.get('conditional_rules', '').strip()}
        text = next(iter(sources)) if len(sources) == 1 else None
        try: when = _routing_conditions([text], fields) if text else []
        except ValueError: when = []
        layouts.append({'id': name, 'name': name, 'field_ids': [f['id'] for f in fields],
                        'record_length': spec.get('options', {}).get('record_length') or max((f.get('rules', {}).get('end') or 0 for f in fields), default=0),
                        'when': when, 'source_condition': text})
    return layouts


def _padding_character(value):
    return ' ' if value.strip().upper() in ('SPACE', 'SPACE (0X20)', '0X20') else value


def normalize_spec(draft):
    if not isinstance(draft, dict): raise SpecError('Specification must be an object')
    spec = copy.deepcopy(draft)
    fields = spec.setdefault('fields', [])
    if not isinstance(fields, list) or any(not isinstance(f, dict) for f in fields):
        raise SpecError('Specification fields must be a list of field objects')
    for index, field in enumerate(fields, 1):
        field.setdefault('id', f'field-{index}')
        for key in COLUMN_KEYS: field[key] = '' if field.get(key) is None else str(field.get(key, ''))
        field['sequence'] = field['sequence'] or str(index)
        old_rules = field.get('rules') if isinstance(field.get('rules'), dict) else {}
        field['mandatory'] = {'Y': 'YES', 'TRUE': 'YES', 'N': 'NO', 'FALSE': 'NO', '?': 'UNKNOWN'}.get(field['mandatory'].strip().upper(), field['mandatory'].strip().upper()) or ('YES' if old_rules.get('required') is True else 'NO' if old_rules.get('required') is False else 'UNKNOWN')
        if not field['source_reference'] and field.get('source'):
            field['source_reference'] = f"{field['source'].get('sheet', '')}!{field['source'].get('row', '')}"
        baseline = infer_rules(field); explicit = field.get('rules', {})
        if not isinstance(explicit, dict): raise SpecError('Field rules must be objects')
        rules = {**baseline, **explicit}
        rules['required'] = {'YES': True, 'NO': False, 'UNKNOWN': None}.get(field['mandatory'])
        if isinstance(explicit.get('padding'), dict): rules['padding'] = {**baseline.get('padding', {}), **explicit['padding']}
        if field['alignment'].upper() in ('LEFT', 'RIGHT', 'EXACT'):
            rules['padding']['side'] = {'LEFT': 'right', 'RIGHT': 'left', 'EXACT': 'none'}[field['alignment'].upper()]
        if field['padding_direction'].upper() in ('TRAILING', 'LEADING', 'RIGHT', 'LEFT', 'NONE'):
            rules['padding']['side'] = {'TRAILING': 'right', 'RIGHT': 'right', 'LEADING': 'left', 'LEFT': 'left', 'NONE': 'none'}[field['padding_direction'].upper()]
        if field['padding_byte']:
            rules['padding']['char'] = _padding_character(field['padding_byte'])
        if field['allowed_values']:
            rules['allowed_values'] = field['allowed_values'].splitlines()
        if field['format'] and rules.get('datatype') == 'date':
            rules['date_format'] = field['format'].upper()
        if 'allow_internal_spaces' in rules: rules['allow_internal_whitespace'] = rules['allow_internal_spaces']
        # The editable source-position columns are authoritative; stale AI rules cannot move them back.
        for key in ('start', 'end', 'length'): rules[key] = baseline[key]
        field['rules'] = rules
    spec.setdefault('options', {})
    spec['layouts'] = _layouts(spec)
    if isinstance(spec['layouts'], list):
        for layout in spec['layouts']:
            if not isinstance(layout, dict): continue
            layout['record_length'] = _integer(layout.get('record_length'))
            if isinstance(layout.get('when', []), list):
                for condition in layout.get('when', []):
                    if isinstance(condition, dict):
                        for key in ('start', 'end'): condition[key] = _integer(condition.get(key))
    spec['issues'] = validate_spec(spec)
    spec['coverage_limitations'] = [issue('UNKNOWN_MANDATORY', 'Mandatory status was not supplied; blank-value validation is unavailable', f['id'], severity='warning') for f in fields if f['mandatory'] == 'UNKNOWN' and f['rules'].get('mapping_type') in ('physical', 'filler')]
    return spec


def validate_spec(spec):
    if not isinstance(spec, dict): return [issue('SPEC_TYPE', 'Specification must be an object')]
    findings, fields = _option_issues(spec.get('options', {})), spec.get('fields', [])
    if not isinstance(fields, list) or not fields or len(fields) > MAX_FIELDS:
        return [issue('FIELD_COUNT', 'An approved specification needs 1–4,000 fields')]
    if any(not isinstance(f, dict) for f in fields): return [issue('FIELD_TYPE', 'Fields must be objects')]
    ids = [f.get('id') for f in fields]
    if any(not isinstance(i, str) or not i for i in ids): return [issue('FIELD_IDS', 'Every field needs a unique stable identifier')]
    if len(set(ids)) != len(ids): findings.append(issue('FIELD_IDS', 'Every field needs a unique stable identifier'))
    sequences = [str(f.get('sequence', '')).strip() for f in fields]
    if any(not re.fullmatch(r'[1-9][0-9]*', value) for value in sequences):
        findings.append(issue('SEQUENCE', 'Every field needs a positive sequence number'))
    layouts = _layouts(spec)
    if not isinstance(layouts, list): return findings + [issue('LAYOUTS', 'Record layouts must be a list')]
    supported_conditions = set()
    by_id = {field['id']: field for field in fields}
    for layout in layouts:
        if not isinstance(layout, dict): continue
        condition_text = layout.get('source_condition')
        references = layout.get('field_ids', [])
        if not isinstance(condition_text, str) or not condition_text or not isinstance(references, list): continue
        if any(not isinstance(field_id, str) for field_id in references): continue
        selected = [by_id[field_id] for field_id in references if field_id in by_id]
        try:
            if layout.get('when') == _routing_conditions([condition_text], selected):
                supported_conditions.update((field['id'], condition_text) for field in selected)
        except ValueError:
            pass
    for f in fields:
        identifier, r = f.get('id'), f.get('rules') or infer_rules(f)
        if not isinstance(r, dict): return findings + [issue('RULE_TYPE', 'Field rules must be objects', identifier)]
        if not str(f.get('field_name', '')).strip(): findings.append(issue('FIELD_NAME', 'Field name is required', identifier))
        if f.get('alignment', '').upper() not in ('', 'LEFT', 'RIGHT', 'EXACT', 'UNSPECIFIED'):
            findings.append(issue('ALIGNMENT', 'Alignment must be LEFT, RIGHT, EXACT or UNSPECIFIED; center alignment is not supported', identifier))
        if f.get('padding_direction', '').upper() not in ('', 'LEADING', 'TRAILING', 'LEFT', 'RIGHT', 'NONE'):
            findings.append(issue('PADDING_DIRECTION', 'Padding direction must be LEADING, TRAILING, LEFT, RIGHT or NONE', identifier))
        if f.get('padding_byte') and len(_padding_character(f['padding_byte']).encode('utf-8')) != 1:
            findings.append(issue('PADDING_BYTE', 'Padding byte must occupy exactly one byte', identifier))
        if f.get('format') and r.get('datatype') != 'date':
            findings.append(issue('UNSUPPORTED_FORMAT', 'Only explicit supported date masks can be executed locally', identifier))
        condition_text = f.get('conditional_rules', '').strip()
        if condition_text and (identifier, condition_text) not in supported_conditions:
            findings.append(issue('UNSUPPORTED_CONDITION', 'Conditional field rules require a supported executable definition', identifier))
        kind = r.get('mapping_type')
        if kind not in ('physical', 'filler', 'constant', 'system'): findings.append(issue('MAPPING_TYPE', 'Resolve the mapping type', identifier)); continue
        source_kind = infer_rules(f).get('mapping_type')
        if source_kind in ('physical', 'filler') and kind in ('system', 'constant'):
            findings.append(issue('SOURCE_MAPPING', 'Physical source mappings cannot disappear through a rule override; correct the source mapping explicitly', identifier))
        if kind == 'system':
            if any(r.get(k) is not None for k in ('start', 'end', 'length')): findings.append(issue('GENERATED_POSITION', 'System mappings must not consume DAT bytes', identifier))
            if r.get('conversion') != 'database': findings.append(issue('SYSTEM_CONVERSION', 'System expressions are database-only metadata, never locally executed', identifier))
            continue
        if kind == 'constant' and r.get('constant') is None: findings.append(issue('CONSTANT', 'Supply an explicit literal constant; a sample is not a constant', identifier))
        physical = kind in ('physical', 'filler') or any(r.get(k) is not None for k in ('start', 'end', 'length'))
        if not physical: continue
        start, end, length = (_integer(r.get(k)) for k in ('start', 'end', 'length'))
        if not start or not end or not length or not 1 <= start <= end <= MAX_RECORD_WIDTH or end - start + 1 != length:
            findings.append(issue('POSITIONS', 'Start, inclusive end and length must agree; positions are never renumbered automatically', identifier))
        if f.get('mandatory') not in ('YES', 'NO', 'UNKNOWN'):
            findings.append(issue('MANDATORY', 'Mandatory must be YES, NO or UNKNOWN', identifier))
        elif f['mandatory'] == 'UNKNOWN':
            findings.append(issue('REQUIRED_RULE', 'Mandatory status is unknown; acknowledge the coverage limit or set YES/NO', identifier,
                                  severity='warning' if spec.get('options', {}).get('acknowledge_unknown_mandatory') is True else 'error'))
        source_rule = _type_rule(str(f.get('datatype_conversion', '')))
        if kind != 'constant' and source_rule.get('unsupported_conversion'):
            findings.append(issue('UNSUPPORTED_CONVERSION', 'Resolve unsupported source conversion: ' + source_rule['unsupported_conversion'], identifier))
        elif kind != 'constant' and (any(r.get(key) != source_rule.get(key) for key in ('datatype', 'conversion')) or
                                     (source_rule.get('datatype') == 'date' and r.get('date_format') != source_rule.get('date_format'))):
            findings.append(issue('SOURCE_DATATYPE', 'Interpreted datatype/conversion must retain every supported source check; edit the declaration explicitly to change it', identifier))
        if r.get('datatype') not in ('char', 'integer', 'decimal', 'float', 'date'):
            findings.append(issue('DATATYPE', 'Resolve the supported datatype check', identifier))
        if r.get('conversion') not in ('identity', 'trim', 'rtrim', 'ltrim', 'to_number', 'to_date'):
            findings.append(issue('CONVERSION', 'Conversion is unsupported; generated code or SQL cannot be executed', identifier))
        if r.get('datatype') == 'date' and r.get('date_format') not in DATE_FORMATS:
            findings.append(issue('DATE_FORMAT', 'Resolve an explicitly supported date mask', identifier))
        padding = r.get('padding', {})
        if not isinstance(padding, dict) or padding.get('side') not in ('left', 'right', 'none') or not isinstance(padding.get('char'), str) or len(padding.get('char')) != 1:
            findings.append(issue('PADDING', 'Resolve padding side and one-byte padding character', identifier))
        elif padding['char'] != ' ' and (not padding['char'].isascii() or ord(padding['char']) < 32 or ord(padding['char']) == 127):
            findings.append(issue('PADDING', 'Padding must be one printable ASCII byte', identifier))
        if r.get('allowed_values') is not None and (not isinstance(r['allowed_values'], list) or not all(isinstance(x, str) for x in r['allowed_values'])):
            findings.append(issue('ALLOWED_VALUES', 'Permitted values must be a list of exact strings', identifier))
        for key in ('allow_internal_spaces', 'allow_internal_whitespace', 'shared_source'):
            if key in r and not isinstance(r[key], bool): findings.append(issue('RULE_VALUE', f'{key} must be true or false', identifier))
    layouts = _layouts(spec)
    if not isinstance(layouts, list) or not layouts: return findings + [issue('LAYOUTS', 'Resolve at least one record layout')]
    seen_layouts, covered = set(), set()
    known_layouts = {layout.get('id') for layout in layouts if isinstance(layout, dict)}
    for field in fields:
        if field.get('record_type') and field['record_type'] not in known_layouts:
            findings.append(issue('RECORD_TYPE', 'Record type must name an explicit layout ID', field['id']))
    by_id = {f.get('id'): f for f in fields}
    for layout in layouts:
        if not isinstance(layout, dict): findings.append(issue('LAYOUT_TYPE', 'Each layout must be an object')); continue
        identifier = layout.get('id')
        if not isinstance(identifier, str) or not identifier or identifier in seen_layouts: findings.append(issue('LAYOUT_IDS', 'Layouts need unique identifiers'))
        seen_layouts.add(identifier)
        references = layout.get('field_ids', [])
        if not isinstance(references, list) or not references or any(not isinstance(x, str) or x not in by_id for x in references): findings.append(issue('LAYOUT_FIELDS', 'Each layout must reference known fields')); continue
        covered.update(references)
        width = _integer(layout.get('record_length'))
        if not width or width > MAX_RECORD_WIDTH: findings.append(issue('RECORD_LENGTH', 'Resolve a positive record length no greater than 1,000,000 bytes'))
        spans = []
        for reference in references:
            r = by_id[reference].get('rules', {})
            if r.get('start') and r.get('end'):
                if width and r['end'] > width: findings.append(issue('RECORD_LENGTH', 'Field extends beyond its record layout', reference))
                for other, previous in spans:
                    overlap = max(r['start'], previous['start']) <= min(r['end'], previous['end'])
                    identical = (r['start'], r['end']) == (previous['start'], previous['end'])
                    if overlap and not identical and not (r.get('shared_source') and previous.get('shared_source')):
                        findings.append(issue('OVERLAP', f'Overlapping source slices with {other} need explicit shared-source approval', reference))
                spans.append((reference, r))
        conditions = layout.get('when', [])
        if not isinstance(conditions, list): findings.append(issue('ROUTING', 'Routing must be an explicit list of byte equality conditions')); continue
        valid_conditions = []
        for condition in conditions:
            if not isinstance(condition, dict): findings.append(issue('ROUTING', 'Routing conditions must be objects')); continue
            start, end = _integer(condition.get('start')), _integer(condition.get('end'))
            if not start or not end or start > end or not isinstance(condition.get('equals'), str) or set(condition) - {'start', 'end', 'equals', 'trim'}:
                findings.append(issue('ROUTING', 'Unsupported or incomplete record routing condition'))
            elif width and end > width: findings.append(issue('ROUTING', 'Record routing reads outside its declared layout'))
            elif 'trim' in condition and not isinstance(condition['trim'], bool): findings.append(issue('ROUTING', 'Routing trim must be explicitly true or false'))
            else: valid_conditions.append({**condition, 'start': start, 'end': end})
        invalid_fields = {finding.get('field_id') for finding in findings if finding.get('severity', 'error') == 'error'}
        encoding = spec.get('options', {}).get('encoding', 'ascii')
        if not isinstance(encoding, str) or encoding.lower() not in ('ascii', 'latin1', 'iso-8859-1', 'cp1252', 'utf-8'): continue
        for reference in references:
            r = by_id[reference].get('rules', {})
            values = r.get('allowed_values')
            if reference in invalid_fields or not isinstance(values, list) or not r.get('start') or not r.get('end'): continue
            local = [{**condition, 'start': condition['start'] - r['start'] + 1,
                      'end': condition['end'] - r['start'] + 1} for condition in valid_conditions
                     if r['start'] <= condition['start'] <= condition['end'] <= r['end']]
            if not local: continue
            possible = False
            for value in [*values, *([''] if r.get('required') is not True else [])]:
                try:
                    data, _ = _format(value, r, encoding)
                    normalized = data.decode(encoding).strip(' ')
                except SpecError: continue
                if (normalized in values or (not normalized and not r.get('required'))) and _route({'when': local}, data, encoding):
                    possible = True; break
            if not possible:
                expected = ', '.join(repr(condition['equals']) for condition in local)[:160]
                findings.append(issue('ROUTING_ALLOWED_VALUES',
                                      f'Allowed values {repr(values)[:160]} cannot satisfy routing value(s) {expected} at this field; revise the specification', reference))
    if set(ids) - covered: findings.append(issue('FIELD_COVERAGE', 'Every source field must belong to an explicitly reviewed layout'))
    for finding in spec.get('issues', []):
        if finding.get('code') in ('FORMULA', 'UNMAPPED_ROW', 'DUPLICATE_COLUMN', 'NO_FIELDS') and finding not in findings: findings.append(finding)
    return findings


def _option_issues(options):
    if not isinstance(options, dict): return [issue('OPTIONS', 'File options must be an object')]
    findings = []
    encoding = options.get('encoding', 'ascii')
    if not isinstance(encoding, str) or encoding.lower() not in ('ascii', 'latin1', 'iso-8859-1', 'cp1252', 'utf-8'):
        findings.append(issue('ENCODING', 'Choose an explicit supported encoding: ascii, latin1, cp1252 or utf-8'))
    if options.get('record_format', 'stream') not in ('stream', 'fixed'): findings.append(issue('RECORD_FORMAT', 'Record format must be stream or fixed'))
    length = options.get('record_length')
    if (length is not None and (not _integer(length) or _integer(length) > MAX_RECORD_WIDTH)) or (options.get('record_format') == 'fixed' and length is None):
        findings.append(issue('RECORD_LENGTH', 'Fixed records need a positive byte length no greater than 1,000,000'))
    if options.get('unknown_records', 'block') not in ('block', 'preserve'): findings.append(issue('UNKNOWN_RECORD_POLICY', 'Unknown records require block or explicit preserve handling'))
    line_ending = options.get('line_ending', 'preserve')
    if line_ending not in ('preserve', 'CRLF', 'LF', 'CR', 'none'): findings.append(issue('LINE_ENDING', 'Line ending must be preserve, CRLF, LF, CR or none'))
    if options.get('record_format') == 'fixed' and line_ending not in ('preserve', 'none'): findings.append(issue('LINE_ENDING', 'Fixed records cannot require stream record separators'))
    if 'preserve_extra_bytes' in options and not isinstance(options['preserve_extra_bytes'], bool): findings.append(issue('EXTRA_BYTES_POLICY', 'Extra-byte preservation must be explicitly true or false'))
    if 'acknowledge_unknown_mandatory' in options and not isinstance(options['acknowledge_unknown_mandatory'], bool):
        findings.append(issue('MANDATORY_ACKNOWLEDGMENT', 'Unknown mandatory coverage acknowledgment must be true or false'))
    for key in ('confirmed_records', 'suspect_records'):
        value = options.get(key, [])
        if not isinstance(value, list) or len(value) > 20000 or any(not isinstance(n, int) or isinstance(n, bool) or not 1 <= n <= 20000 for n in value):
            findings.append(issue('RECORD_SELECTION', f'{key} must contain valid physical record numbers'))
    for key in ('detail_record', 'record_number'):
        if key in options and (not isinstance(options[key], int) or isinstance(options[key], bool) or not 1 <= options[key] <= 20000):
            findings.append(issue('RECORD_SELECTION', f'{key} must be a positive record number'))
    return findings


def _settings(spec, options):
    if not isinstance(spec, dict) or not isinstance(spec.get('options', {}), dict) or (options is not None and not isinstance(options, dict)):
        raise SpecError('Specification and file options must be objects')
    settings = {**spec.get('options', {}), **(options or {})}
    settings.setdefault('encoding', 'ascii'); settings.setdefault('record_format', 'stream')
    if 'record_length' in settings: settings['record_length'] = _integer(settings['record_length'])
    return settings


def _route(layout, row, encoding):
    unknown = False
    for condition in layout.get('when', []):
        if len(row) < condition['end']: unknown = True; continue
        try: text = row[condition['start'] - 1:condition['end']].decode(encoding)
        except UnicodeDecodeError: return None
        if condition.get('trim', False): text = text.strip(' ')
        if text != condition['equals']: return False
    return None if unknown else True


def _edits(edits):
    result = {}
    if isinstance(edits, dict):
        for key, value in edits.items():
            if not isinstance(key, tuple) or len(key) != 2 or not isinstance(key[0], int) or isinstance(key[0], bool) or key[0] < 1 or not isinstance(key[1], str):
                raise SpecError('Edit keys must be (positive record number, stable field_id) pairs')
            result[key] = value
        return result
    if edits is not None and (not isinstance(edits, list) or len(edits) > 500): raise SpecError('An edit request supports at most 500 field changes')
    for edit in edits or []:
        if not isinstance(edit, dict): raise SpecError('Edits must be field change objects')
        key = (edit.get('record'), edit.get('field_id', edit.get('field')))
        if not isinstance(key[0], int) or isinstance(key[0], bool) or key[0] < 1 or not isinstance(key[1], str):
            raise SpecError('Each edit needs a positive record number and stable field_id')
        if key in result: raise SpecError('Duplicate field edit in the same request')
        result[key] = edit.get('value')
    return result


def _format(value, rules, encoding):
    if not isinstance(value, str): raise SpecError('Edited values must be strings; use an empty string for blank')
    value = value.strip(' ')
    if any(ord(char) < 32 or ord(char) == 127 for char in value): raise SpecError('Values cannot contain record delimiters or control characters')
    if rules.get('allow_internal_spaces', rules.get('allow_internal_whitespace')) is False and any(c.isspace() for c in value): raise SpecError('Internal whitespace is prohibited by this field rule')
    try: encoded = value.encode(encoding, 'strict')
    except UnicodeEncodeError: raise SpecError('Value cannot be encoded using the approved encoding') from None
    width, pad = rules['length'], rules['padding']
    if len(encoded) > width: raise SpecError(f'Value exceeds its {width}-byte field; it will not be truncated')
    if pad['side'] == 'none':
        if len(encoded) != width: raise SpecError(f'Exact-width field requires {width} bytes')
        return encoded, value
    padding = pad['char'].encode(encoding)
    if pad['side'] == 'right': return encoded.ljust(width, padding), value
    if padding == b'0' and encoded[:1] in (b'+', b'-'):
        return encoded[:1] + encoded[1:].rjust(width - 1, padding), value
    return encoded.rjust(width, padding), value


def _field_checks(field, data, encoding):
    r, checks = field['rules'], []
    try: text = data.decode(encoding)
    except UnicodeDecodeError:
        return {'value': '', 'checks': [{'code': 'ENCODING', 'passed': False,
                'message': 'Field byte boundaries split or contain an invalid encoded character'}],
                'passed': 0, 'total': 1, 'percentage': 0.0, 'status': 'failed',
                'staging_preview': None, 'preview_label': 'Local preview; not an Oracle database result'}
    value = text.strip(' ')
    def check(code, passed, message, failure): checks.append({'code': code, 'passed': passed, 'message': failure if passed is False else message})
    check('WIDTH', len(data) == r['length'], f'Field occupies {r["length"]} bytes', f'Field has {len(data)} bytes; expected {r["length"]}')
    if r.get('required') is None: check('MANDATORY_UNVERIFIED', None, 'Mandatory status is unknown; blank-value rule cannot be verified', '')
    else: check('REQUIRED', bool(value) or not r.get('required'), 'Required value is present' if r.get('required') else 'Blank value is permitted', 'Required value is blank')
    check('CONTROL_CHARACTERS', all(ord(c) >= 32 and ord(c) != 127 for c in text), 'No prohibited control characters', 'Field contains prohibited control characters')
    if r.get('allow_internal_spaces', r.get('allow_internal_whitespace')) is False: check('INTERNAL_WHITESPACE', not any(c.isspace() for c in value), 'No prohibited internal whitespace', 'Field contains prohibited internal whitespace')
    try: formatted, _ = _format(value, r, encoding); padding_ok = formatted == data
    except SpecError: padding_ok = False
    check('PADDING', padding_ok, 'Padding follows the approved rule', 'Padding does not match the approved rule')
    if r.get('constant') is not None: check('CONSTANT', value == r['constant'], 'Value equals its declared constant', 'Value differs from its declared constant')
    if r.get('allowed_values') is not None: check('ALLOWED_VALUES', value in r['allowed_values'] or (not value and not r['required']), 'Value belongs to the approved set', f'Value is not in the allowed set {repr(r["allowed_values"])[:160]}')
    datatype = r['datatype']
    if value and datatype != 'char':
        valid = True
        if datatype == 'integer': valid = bool(re.fullmatch(r'[+-]?[0-9]+', value))
        elif datatype in ('decimal', 'float'):
            expression = r'[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)' + (r'(?:[eE][+-]?[0-9]+)?' if datatype == 'float' else '')
            valid = bool(re.fullmatch(expression, value))
            try: valid = valid and Decimal(value).is_finite()
            except InvalidOperation: valid = False
        elif datatype == 'date':
            try:
                mask = r['date_format']; valid = len(value) == len(mask)
                datetime.strptime(value, DATE_FORMATS[mask])
            except (ValueError, KeyError): valid = False
        check('DATATYPE', valid, f'Value matches the approved {datatype} format', f'Value does not match the approved {datatype} format')
    if value and r.get('conversion') == 'to_number' and datatype == 'char':
        try: valid = Decimal(value).is_finite()
        except InvalidOperation: valid = False
        check('CONVERSION', valid, 'Number-conversion preview is evaluable', 'Value cannot be converted to a number')
    passed = sum(c['passed'] is True for c in checks)
    preview = value
    if any(c['passed'] is not True for c in checks): preview = None
    elif not value: preview = None
    elif r.get('conversion') == 'to_number': preview = str(Decimal(value))
    elif r.get('conversion') == 'to_date': preview = datetime.strptime(value, DATE_FORMATS[r['date_format']]).date().isoformat()
    elif r.get('conversion') == 'rtrim': preview = text.rstrip(' ')
    elif r.get('conversion') == 'ltrim': preview = text.lstrip(' ')
    return {'value': value, 'checks': checks, 'passed': passed, 'total': len(checks),
            'percentage': round(100 * passed / len(checks), 2) if checks else None,
            'status': 'passed' if all(c['passed'] for c in checks) else 'failed', 'staging_preview': preview,
            'preview_label': 'Local preview; not an Oracle database result'}


def _process(spec, raw, edits=None, options=None):
    if not isinstance(raw, bytes): raise SpecError('DAT must contain original byte data')
    spec = normalize_spec(spec)
    settings, findings = _settings(spec, options), validate_spec(spec)
    findings.extend(_option_issues(settings))
    if any(f.get('severity', 'error') == 'error' for f in findings):
        raise SpecError('Resolve the specification and file options before validating DAT', findings)
    if len(raw) > MAX_BYTES: raise SpecError('DAT exceeds 12 MB')
    try:
        if settings['encoding'].lower() == 'utf-8':
            raw.decode('utf-8', 'strict')
            encoding = 'utf-8'
        else: encoding = check_encoding(raw, settings['encoding'])
        lines = parse_lines(raw, settings['record_format'], settings.get('record_length'))
    except ValueError as exc: raise SpecError(str(exc)) from None
    changes, seen_edits = _edits(edits), set()
    layouts, by_id = _layouts(spec), {f['id']: f for f in spec['fields']}
    records, summaries, output, changed_records = [], [], [], []
    passed = total = pending = blocking = all_issue_count = 0
    detail = settings.get('detail_record', settings.get('record_number'))
    confirmed = set(settings.get('confirmed_records', []))
    for number, (row, ending) in enumerate(lines, 1):
        candidates = [(layout, _route(layout, row, encoding)) for layout in layouts]
        matches = [layout for layout, matches in candidates if matches is True]
        record_issues, fields = [], []
        if len(matches) != 1 or any(state is None for _, state in candidates):
            message = 'Record matches multiple layouts' if len(matches) > 1 else 'Record layout is unknown or cannot be evaluated'
            preserve = not matches and settings.get('unknown_records') == 'preserve' and not any(key[0] == number for key in changes)
            finding = issue('AMBIGUOUS_LAYOUT' if len(matches) > 1 else 'UNKNOWN_LAYOUT', message, record=number, severity='warning' if preserve else 'error')
            record_issues.append(finding)
            if len(findings) < 1000: findings.append(finding)
            all_issue_count += 1
            status = 'preserved' if preserve else 'needs_review'
            summaries.append({'number': number, 'layout_id': None, 'status': status, 'percentage': None, 'issues': record_issues, 'issue_count': 1})
            if detail is None or detail == number: records.append({'number': number, 'layout_id': None, 'fields': [], 'issues': record_issues, 'status': status})
            output.append(row + ending); pending += 1; blocking += 1
            continue
        layout = matches[0]; width = layout['record_length']
        from tasaa import analyze as analyze_alignment, verify_preservation
        alignment = analyze_alignment(spec, layout, row, encoding=encoding)
        repair_mapping = (alignment.get('alternatives') or [{}])[0].get('mapping')
        preservation = (verify_preservation(row, alignment['candidate'], repair_mapping)
                        if isinstance(alignment.get('candidate'), bytes) and repair_mapping else None)
        certified_repair = (alignment.get('status') == 'certified'
                            and ((alignment.get('certification') or {}).get('level') in (
                                'GLOBALLY_CERTIFIED_WITHIN_FORMAL_CONSTRAINT_MODEL',
                                'CERTIFIED_SAFE_REPAIR')
                                or alignment.get('classification') == 'CERTIFIED_SAFE_REPAIR')
                            and alignment.get('validation_status') == 'PASSED'
                            and not alignment.get('resource_limit_reached')
                            and preservation == alignment.get('provenance')
                            and isinstance(alignment.get('candidate'), bytes))
        if certified_repair:
            from validation_pipeline import compile_pipeline as _comp_pipe, validate_record as _val_rec
            try:
                _pipe = _comp_pipe(spec)
                _val = _val_rec(_pipe, alignment['candidate'], layout['id'])
                if not _val['valid']:
                    certified_repair = False
            except Exception:
                certified_repair = False
        source_row = alignment['candidate'] if certified_repair else row
        candidate = bytearray(source_row)
        suspect = len(row) != width or number in settings.get('suspect_records', [])
        desired_ending = {'CRLF': b'\r\n', 'LF': b'\n', 'CR': b'\r', 'none': b''}.get(settings.get('line_ending'))
        if desired_ending is not None and ending != desired_ending:
            record_issues.append(issue('LINE_ENDING', 'Record ending differs from the approved file rule; original ending was preserved', record=number))
        if suspect and number not in confirmed and not certified_repair:
            record_issues.append(issue('UNCERTAIN_BOUNDARIES', 'Record width or field boundaries need explicit confirmation before formatting', record=number))
            pending += 1
        if len(row) > width and not certified_repair and not settings.get('preserve_extra_bytes', False):
            extra_value = next(((offset, byte) for offset, byte in enumerate(row[width:], width + 1) if byte != 32), None)
            extra_note = (f'; extra non-space byte at position {extra_value[0]} (0x{extra_value[1]:02X}) requires an approved owner'
                          if extra_value else '; extra spaces require approved padding or filler ownership')
            record_issues.append(issue('RECORD_LENGTH', f'Record has {len(row)} bytes; approved layout has {width}{extra_note}; extra bytes are never truncated', record=number))
        writes = {}
        for identifier in layout['field_ids']:
            field, r = by_id[identifier], by_id[identifier]['rules']
            key = (number, identifier); editing = key in changes
            if r['mapping_type'] == 'system' or (r['mapping_type'] == 'constant' and not r.get('start')):
                if editing: record_issues.append(issue('READ_ONLY', 'Generated and nonphysical constant mappings are not DAT cells', identifier, number)); seen_edits.add(key)
                fields.append({'id': identifier, **{k: field.get(k, '') for k in COLUMN_KEYS}, 'original_value': None,
                               'value': r.get('constant'), 'checks': [], 'original_checks': [], 'passed': 0, 'total': 0,
                               'percentage': None, 'original_percentage': None, 'status': 'not_applicable',
                               'staging_preview': r.get('constant'), 'preview_label': 'Database-generated value is not locally evaluated', 'issues': []})
                continue
            original = source_row[r['start'] - 1:r['end']]
            before = _field_checks(field, original, encoding)
            value = changes[key] if editing else before['value']
            field_issues = []
            if any(check['code'] == 'ENCODING' and check['passed'] is False for check in before['checks']):
                field_issues.append(issue('ENCODING', 'Original field bytes split or contain an invalid encoded character', identifier, number))
            if editing: seen_edits.add(key)
            if r['mapping_type'] == 'filler':
                if editing: field_issues.append(issue('READ_ONLY', 'Filler bytes are managed by the approved spec', identifier, number))
                value = r.get('filler_value', '')
            formatted = original
            try:
                formatted, normalized_value = _format(value, r, encoding)
                for offset, byte in enumerate(formatted, r['start'] - 1):
                    if offset in writes and writes[offset] != byte:
                        field_issues.append(issue('CONFLICTING_WRITES', 'Shared source mappings propose different bytes', identifier, number)); break
                    writes[offset] = byte
                if not field_issues:
                    if len(candidate) < r['end']: candidate.extend(b' ' * (r['end'] - len(candidate)))
                    candidate[r['start'] - 1:r['end']] = formatted
                after = _field_checks(field, formatted, encoding)
            except SpecError as exc:
                normalized_value = value
                if isinstance(value, str):
                    try: after = _field_checks(field, value.strip(' ').encode(encoding, 'strict'), encoding)
                    except UnicodeEncodeError: after = {**before, 'checks': [*before['checks'], {'code': 'ENCODING', 'passed': False, 'message': 'Value cannot be encoded'}]}
                else: after = {**before, 'checks': [*before['checks'], {'code': 'VALUE_TYPE', 'passed': False, 'message': 'Values must be strings'}]}
                after['passed'] = sum(c['passed'] is True for c in after['checks']); after['total'] = len(after['checks'])
                after['percentage'] = round(100 * after['passed'] / after['total'], 2); after['status'] = 'failed'
                field_issues.append(issue('VALUE_FORMAT', str(exc), identifier, number))
            for check in after['checks']:
                if check['passed'] is not True: field_issues.append(issue(check['code'], check['message'], identifier, number))
            if formatted != original and not editing and number not in confirmed:
                field_issues.append(issue('UNAPPROVED_REFORMAT', 'Formatting would change original bytes; review and confirm this record before export', identifier, number))
            if suspect and number not in confirmed:
                after = {**after, 'status': 'needs_review', 'percentage': None}
            if field_issues and after['status'] == 'passed': after = {**after, 'status': 'failed', 'percentage': None}
            fields.append({'id': identifier, **{k: field.get(k, '') for k in COLUMN_KEYS}, **after,
                           'original_value': before['value'], 'original_bytes_hex': original.hex(),
                           'value': normalized_value, 'original_checks': before['checks'], 'original_percentage': before['percentage'],
                           'modified': formatted != original, 'issues': field_issues})
            record_issues.extend(field_issues); passed += after['passed']; total += after['total']
        implicit_change = any(field.get('modified') and (number, field['id']) not in changes for field in fields)
        if implicit_change:
            record_issues.append(issue('UNVERIFIED_STRUCTURAL_CHANGE',
                                       'Implicit byte changes require a TASAA-certified repair; edit values explicitly or review the mapping',
                                       record=number))
        if len(candidate) != width and not (len(candidate) > width and settings.get('preserve_extra_bytes', False)):
            record_issues.append(issue('RECORD_LENGTH', 'Candidate does not occupy the approved record length', record=number))
        if len(candidate) > len(row) and not certified_repair and any(offset not in writes for offset in range(len(row), len(candidate))):
            record_issues.append(issue('MISSING_UNMAPPED_BYTES', 'Missing unmapped bytes cannot be invented as padding; approve a field/filler rule or supply the source bytes', record=number))
        changed = bytes(candidate) != row
        if changed: changed_records.append(number)
        # Validate the exact bytes after all writes, including aliases and routing markers.
        for field in fields:
            r = by_id[field['id']]['rules']
            if not r.get('start'): continue
            exact = _field_checks(by_id[field['id']], bytes(candidate[r['start'] - 1:r['end']]), encoding)
            if exact['status'] != 'passed' and not field['issues']:
                record_issues.append(issue('CANDIDATE_BYTES', 'Exact serialized bytes fail field checks', field['id'], number))
        if _route(layout, bytes(candidate), encoding) is not True or sum(_route(x, bytes(candidate), encoding) is True for x in layouts) != 1:
            record_issues.append(issue('ROUTING_CHANGED', 'Formatting or edits changed record routing; review the layout explicitly', record=number))
        all_issue_count += len(record_issues)
        findings.extend(record_issues[:max(0, 1000 - len(findings))]); blocking += sum(x['severity'] == 'error' for x in record_issues)
        record_passed = sum(f['passed'] for f in fields); record_total = sum(f['total'] for f in fields)
        status = 'needs_review' if any(x['code'] == 'UNCERTAIN_BOUNDARIES' for x in record_issues) else 'failed' if record_issues else 'passed'
        percentage = round(100 * record_passed / record_total, 2) if record_total and not any(x['code'] == 'UNCERTAIN_BOUNDARIES' for x in record_issues) else None
        structural_issues = [x for x in record_issues if x['code'] in ('RECORD_LENGTH', 'UNCERTAIN_BOUNDARIES', 'MISSING_UNMAPPED_BYTES', 'CONFLICTING_WRITES', 'UNVERIFIED_STRUCTURAL_CHANGE')]
        structural_compliance = len(structural_issues) == 0 and (len(candidate) == width or certified_repair)
        data_rule_compliance = not any(x['severity'] == 'error' for x in record_issues if x not in structural_issues)
        summary = {'number': number, 'layout_id': layout['id'], 'status': status, 'percentage': percentage,
                   'issues': record_issues[:5], 'issue_count': len(record_issues), 'changed': changed,
                   'structural_compliance': structural_compliance, 'data_rule_compliance': data_rule_compliance}
        summaries.append(summary)
        if detail is None or detail == number:
            alignment = alignment or {}
            alignment = {k: v for k, v in alignment.items() if k not in ('candidate', 'alternatives')}
            records.append({**summary, 'fields': fields, 'original_length': len(row), 'candidate_length': len(candidate),
                            'alignment': alignment})
        output.append(bytes(candidate) + ending)
    if set(changes) - seen_edits: raise SpecError('An edit references a missing record, unknown field or another layout')
    if not lines: findings.append(issue('EMPTY_DAT', 'DAT contains no records')); blocking += 1; all_issue_count += 1
    candidate = b''.join(output)
    report = {'records': records, 'record_summaries': summaries, 'issues': findings, 'issue_count': all_issue_count, 'valid': blocking == 0,
              'candidate_checksum': hashlib.sha256(candidate).hexdigest(), 'original_checksum': hashlib.sha256(raw).hexdigest(),
              'changed_records': changed_records, 'summary': {'passed': passed, 'total': total,
              'percentage': round(100 * passed / total, 2) if total and not pending else None,
              'pending_count': pending, 'blocking_count': blocking, 'record_count': len(lines), 'changed_records': len(changed_records),
              'structural_compliance': all(s.get('structural_compliance', False) for s in summaries),
              'data_rule_compliance': all(s.get('data_rule_compliance', False) for s in summaries)},
              'encoding': encoding, 'record_format': settings['record_format'], 'preview_label': 'Local preview; not an Oracle database result'}
    return report, candidate


def inspect_dat(spec, raw, edits=None, options=None):
    return _process(spec, raw, edits, options)[0]


def render_dat(spec, raw, edits=None, options=None):
    report, candidate = _process(spec, raw, edits, options)
    if not report['valid']: raise SpecError('Candidate DAT still has unresolved or failed checks', report['issues'])
    from validation_pipeline import compile_pipeline, validate_record
    pipeline = compile_pipeline(spec)
    settings = _settings(spec, options)
    lines = parse_lines(candidate, settings['record_format'], settings.get('record_length'))
    for number, (row, _) in enumerate(lines, 1):
        result = validate_record(pipeline, row)
        if not result['valid']:
            raise SpecError('Compiled validation failed on the exact export bytes',
                            [issue(check['code'], 'Compiled rule failed on serialized DAT', check.get('field_id'), number)
                             for check in result['issues']])
    return candidate
