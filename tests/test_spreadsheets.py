import io
import pytest
from openpyxl import Workbook
from spreadsheets import import_spreadsheet,default_mapping,reconcile

CTL='LOAD DATA INTO TABLE T (ID POSITION(1:4) CHAR, AMOUNT POSITION(5:10) DECIMAL EXTERNAL)'


def test_csv_exact_decimals_zeroes_blank_and_duplicate_keys():
    sh=import_spreadsheet('ref.csv',b'ID,AMOUNT\n0001,10.230\n0002,\n0003,1.00\n0003,2.00\n')
    sheet=sh['sheets'][0]
    assert default_mapping(CTL,sheet)=={'ID':'ID','AMOUNT':'AMOUNT'}
    report=reconcile(CTL,b'000110.200\n0002      \n00031.00  \n0004      ',sh,{'ID':'ID','AMOUNT':'AMOUNT'},['ID'],'CSV')
    assert report['conflicts'][0]['value']=='10.230'
    assert report['conflicts'][0]['original_value']=='10.200'
    assert len(report['duplicate_keys'])==1
    assert report['missing_supplement_records']==[3,4]
    assert not any(e['field']=='AMOUNT' and e['record']==2 for e in report['evidence'])


def test_xlsx_multiple_sheets_preamble_zero_format_and_formula():
    wb=Workbook();ws=wb.active;ws.title='Transactions'
    ws.append(['Monthly reference']);ws.append(['ID','AMOUNT']);ws.append([1,'123.4500']);ws['A3'].number_format='0000'
    ws.append(['0002','=1+2'])
    other=wb.create_sheet('Mapping');other.append(['FIELD_NAME','START_POS','END_POS']);other.append(['ID',1,4])
    buf=io.BytesIO();wb.save(buf)
    result=import_spreadsheet('ref.xlsx',buf.getvalue())
    assert len(result['sheets'])==2
    assert result['sheets'][0]['header_row']==2
    assert result['sheets'][0]['rows'][0]['values']=={'ID':'0001','AMOUNT':'123.4500'}
    assert result['sheets'][0]['rows'][1]['values']['AMOUNT']==''
    assert result['warnings']
    assert result['sheets'][1]['kind']=='mapping'


def test_manual_header_mapping_and_exact_keys():
    sh=import_spreadsheet('r.csv',b'Payment Reference,Value\n0001,10.00\n')
    r=reconcile(CTL,b'000110.00 ',sh,{'PAYMENT_REFERENCE':'ID','VALUE':'AMOUNT'},['ID'],'CSV')
    assert r['matched_records']==1 and not r['conflicts']
    with pytest.raises(ValueError):reconcile(CTL,b'000110.00 ',sh,{'VALUE':'AMOUNT'},[],'CSV')
    assert reconcile(CTL,b'1   10.00 ',sh,{'PAYMENT_REFERENCE':'ID','VALUE':'AMOUNT'},['ID'],'CSV')['missing_supplement_records']==[1]


def test_duplicate_dat_keys_cannot_authorize_reconciliation():
    sh=import_spreadsheet('r.csv',b'ID,AMOUNT\n0001,10.00\n')
    r=reconcile(CTL,b'000111.00 \n000112.00 ',sh,{'ID':'ID','AMOUNT':'AMOUNT'},['ID'],'CSV')
    assert r['duplicate_dat_keys'] and not r['evidence'] and not r['conflicts']


def test_field_specification_comparison_is_advisory():
    from spreadsheets import compare_field_specifications
    sh=import_spreadsheet('spec.csv',b'FIELD_NAME,START_POS,END_POS,TYPE\nID,1,4,CHAR\nAMOUNT,5,12,DECIMAL EXTERNAL\n',role='field_specifications')
    r=compare_field_specifications(CTL,sh,'CSV',{'FIELD_NAME':'FIELD_NAME','START_POS':'START','END_POS':'END','TYPE':'DATATYPE'})
    assert r['matched_records']==1 and len(r['spec_conflicts'])==1
    assert not r['evidence'] and not r['conflicts']


def test_excel_display_format_never_truncates_or_rounds_stored_business_value():
    from spreadsheets import cell_text
    assert cell_text(1.999,'0000')=='1.999'
    assert cell_text(1.999,'0.00')=='1.999'
    assert cell_text(1.2,'0.00')=='1.20'
    assert cell_text(1,'0000')=='0001'
